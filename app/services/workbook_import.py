"""One-time, idempotent import of workbook transactions into the ERP database."""
import argparse
import hashlib
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import Base, SessionLocal, engine
from app.models import Category, DataImportBatch, Investment, Operation, ReportGroup
from app.seed import seed_db

REPORT_GROUPS = {
    "revenue": ("Доход", 10),
    "direct": ("Прямой расход", 20),
    "overhead": ("Накладные расходы", 30),
    "taxes": ("Налоги и сборы", 40),
    "other": ("Прочие расходы (Excel)", 50),
}
ACCOUNT_RE = re.compile(r"^\s*(\d+)\s+(.+?)\s*$")
CENT = Decimal("0.01")


@dataclass(frozen=True)
class ImportSummary:
    source_filename: str
    source_sha256: str
    operations_imported: int
    investments_imported: int
    duplicate_rows_skipped: int


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _account_parts(value, sheet_name: str, row_number: int) -> tuple[str, str]:
    account = str(value or "").strip()
    if account.casefold() in {"(пусто)", "пусто"}:
        account = ""
    match = ACCOUNT_RE.match(account)
    if match:
        return match.group(1), match.group(2).strip()
    return "", account


def _report_group(db: Session, section: str) -> ReportGroup:
    name, sort_order = REPORT_GROUPS[section]
    group = (
        db.query(ReportGroup)
        .filter(ReportGroup.name == name, ReportGroup.section == section)
        .first()
    )
    if group is None:
        group = ReportGroup(name=name, section=section, sort_order=sort_order)
        db.add(group)
        db.flush()
    return group


def _section_for_account(code: str) -> str:
    if code.startswith("699"):
        return "taxes"
    if code.startswith(("20", "408", "409")):
        return "direct"
    if code.startswith(("40", "41")):
        return "overhead"
    return "other"


def _category(
    db: Session,
    *,
    name: str,
    kind: str,
    section: str,
) -> Category:
    category = (
        db.query(Category)
        .filter(
            Category.name == name,
            Category.kind == kind,
            Category.parent_id.is_(None),
        )
        .first()
    )
    group = _report_group(db, section)
    if category is None:
        category = Category(
            name=name,
            kind=kind,
            report_group_id=group.id,
        )
        db.add(category)
        db.flush()
    elif category.report_group_id is None:
        category.report_group_id = group.id
    return category


def import_workbook(
    db: Session,
    workbook_bytes: bytes,
    source_filename: str,
    *,
    dry_run: bool = False,
) -> ImportSummary:
    """Import dated transaction rows, skipping exact duplicates across sheets.

    The "Инвестиции" worksheet is stored in the investments table. Other
    positive rows become income operations and negative rows become expenses.
    """
    source_hash = hashlib.sha256(workbook_bytes).hexdigest()
    if db.query(DataImportBatch.id).filter_by(source_sha256=source_hash).first():
        raise ValueError(f"Файл уже был импортирован: {source_filename}")

    try:
        workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
        operations: list[Operation] = []
        investments: list[Investment] = []
        seen_in_other_sheets: dict[tuple, str] = {}
        duplicates_skipped = 0

        for worksheet in workbook.worksheets:
            sheet_name = worksheet.title.strip()
            is_investment_sheet = sheet_name.casefold() == "инвестиции"
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if len(row) < 5 or not any(value is not None for value in row):
                    continue
                transaction_date = _as_date(row[1])
                if transaction_date is None:
                    continue
                try:
                    amount = Decimal(str(row[4])).quantize(CENT, rounding=ROUND_HALF_UP)
                except (InvalidOperation, ValueError, TypeError) as exc:
                    raise ValueError(
                        f"Некорректная сумма на листе «{sheet_name}», строка {row_number}"
                    ) from exc
                if not amount.is_finite() or amount == 0:
                    raise ValueError(
                        f"Нулевая или некорректная сумма на листе «{sheet_name}», строка {row_number}"
                    )

                raw_account = str(row[2] or "").strip()
                comment = str(row[3] or "").strip()
                signature = (
                    transaction_date,
                    " ".join(raw_account.split()).casefold(),
                    " ".join(comment.split()).casefold(),
                    amount,
                )
                previous_sheet = seen_in_other_sheets.get(signature)
                if previous_sheet is not None and previous_sheet != sheet_name:
                    duplicates_skipped += 1
                    continue
                seen_in_other_sheets.setdefault(signature, sheet_name)

                code, account_name = _account_parts(raw_account, sheet_name, row_number)
                source_note = (
                    f"Источник Excel: {sheet_name}; "
                    f"счёт: {raw_account or 'не указан'}"
                )
                full_comment = " | ".join(part for part in (comment, source_note) if part)

                if is_investment_sheet:
                    investments.append(
                        Investment(
                            date=transaction_date,
                            amount=abs(amount),
                            comment=full_comment,
                        )
                    )
                    continue

                kind = "income" if amount > 0 else "expense"
                section = "revenue" if kind == "income" else _section_for_account(code)
                category_name = account_name or sheet_name
                category = _category(
                    db,
                    name=category_name,
                    kind=kind,
                    section=section,
                )
                operations.append(
                    Operation(
                        date=transaction_date,
                        kind=kind,
                        category_id=category.id,
                        amount=abs(amount),
                        comment=full_comment or None,
                    )
                )

        if not operations and not investments:
            raise ValueError("В книге не найдено строк операций с датой и суммой")

        if not dry_run:
            db.add(
                DataImportBatch(
                    source_sha256=source_hash,
                    source_filename=Path(source_filename).name,
                    operations_imported=len(operations),
                    investments_imported=len(investments),
                )
            )
            db.add_all(operations)
            db.add_all(investments)
            db.commit()
        else:
            db.rollback()

        return ImportSummary(
            source_filename=Path(source_filename).name,
            source_sha256=source_hash,
            operations_imported=len(operations),
            investments_imported=len(investments),
            duplicate_rows_skipped=duplicates_skipped,
        )
    except Exception:
        db.rollback()
        raise


def main() -> None:
    parser = argparse.ArgumentParser(description="Import a dated Excel ledger into the ERP database.")
    parser.add_argument("workbook", type=Path, help="Path to an .xlsx workbook")
    parser.add_argument("--dry-run", action="store_true", help="Validate and summarize without importing rows")
    args = parser.parse_args()
    workbook_bytes = args.workbook.read_bytes()

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        if not args.dry_run:
            seed_db(db)
        summary = import_workbook(
            db,
            workbook_bytes,
            args.workbook.name,
            dry_run=args.dry_run,
        )
        print(f"Workbook: {summary.source_filename}")
        print(f"SHA-256: {summary.source_sha256}")
        print(f"Operations: {summary.operations_imported}")
        print(f"Investments: {summary.investments_imported}")
        print(f"Duplicates skipped: {summary.duplicate_rows_skipped}")
        print("Mode: dry run" if args.dry_run else "Import committed")
    finally:
        db.close()


if __name__ == "__main__":
    main()
