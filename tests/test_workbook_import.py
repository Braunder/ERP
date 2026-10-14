from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

import app.database as database_module
from app.models import Category, DataImportBatch, Investment, Operation
from app.services.workbook_import import import_workbook


def _workbook_bytes():
    workbook = Workbook()
    revenue = workbook.active
    revenue.title = "Выручка"
    revenue.append([None, "Дата", "Счёт", "Комментарий", "Сумма"])
    revenue.append([None, datetime(2026, 6, 4), "10101 Выручка основная - проект", "Выручка", 2200])

    expenses = workbook.create_sheet("Расходы")
    expenses.append([None, "Дата", "Счёт", "Комментарий", "Сумма"])
    expenses.append([None, datetime(2026, 6, 8), "40403 Эквайринг", "Эквайринг", -475])

    acquiring = workbook.create_sheet("Эквайринг")
    acquiring.append([None, "Дата", "Счёт", "Комментарий", "Сумма"])
    acquiring.append([None, datetime(2026, 6, 8), "40403 Эквайринг", "Эквайринг", -475])

    investments = workbook.create_sheet("Инвестиции")
    investments.append([None, "Дата", "Счёт", "Комментарий", "Сумма"])
    investments.append([None, datetime(2026, 6, 9), "20399 Прочие расходы подрядчиков - проект", "Материалы", -10000])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_import_workbook_persists_unique_operations_and_investments(client):
    payload = _workbook_bytes()
    db = database_module.SessionLocal()
    try:
        summary = import_workbook(db, payload, "ledger.xlsx")

        assert summary.operations_imported == 2
        assert summary.investments_imported == 1
        assert summary.duplicate_rows_skipped == 1
        assert db.query(Operation).count() == 2
        assert db.query(Investment).count() == 1
        assert db.query(DataImportBatch).count() == 1

        revenue = db.query(Operation).filter_by(kind="income").one()
        assert revenue.amount == 2200
        assert revenue.category.name == "Выручка основная - проект"

        expense = db.query(Operation).filter_by(kind="expense").one()
        assert expense.amount == 475
        assert expense.category.name == "Эквайринг"
        assert "Источник Excel: Расходы" in expense.comment

        investment = db.query(Investment).one()
        assert investment.amount == 10000
        assert "20399 Прочие расходы подрядчиков - проект" in investment.comment

        with pytest.raises(ValueError, match="уже был импортирован"):
            import_workbook(db, payload, "ledger.xlsx")
    finally:
        db.close()


def test_import_workbook_dry_run_does_not_persist_data(client):
    db = database_module.SessionLocal()
    try:
        initial_categories = db.query(Category).count()
        summary = import_workbook(db, _workbook_bytes(), "preview.xlsx", dry_run=True)
        assert summary.operations_imported == 2
        assert summary.investments_imported == 1
        assert db.query(Operation).count() == 0
        assert db.query(Investment).count() == 0
        assert db.query(Category).count() == initial_categories
        assert db.query(DataImportBatch).count() == 0
    finally:
        db.close()
